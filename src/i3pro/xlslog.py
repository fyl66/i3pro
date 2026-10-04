"""读 Excel（``.xlsx``）成场次（ticket #31）。

**为什么读用 openpyxl、写仍然用自己那份**（见 ``xlsx.py`` 与 ADR-0003）：写是我们
自己造的格子（数字、文本、空），格式由我们说了算；读是别人造的格子——共享字符串、
内联字符串、1900 日期序列号、公式的缓存值、多张 sheet、合并单元格——每一样都能
变成一个"看着成功其实错了"的坑。读这一侧不值得自己再踩一遍。

装配成一个**场次**那一步不在这里：表头在哪、单位在哪、哪一列是时间、列名走原名/
别名/手工覆盖、报告怎么写，全部复用 :func:`i3pro.csvlog.session_from_frame`——
Excel 会话与 ``.ld`` / CSV 会话在下游是同一种东西。

读不了的三种工作簿**大声报错**（宏 / 图表 / 外部链接），因为那三种情况下"读到的
数字"与"用户以为的数字"可能不是一回事；公式没有缓存值同样报错。报错都带下一步。
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pandas as pd

from . import csvlog
from .sidecar import SidecarError

__all__ = ["read_xlsx_session", "sheet_names", "UNSUPPORTED_FEATURES"]

#: 表头与单位行只会出现在开头这些行里（和 CSV 那边 ``scan_rows`` 的默认值同一个数字）。
TOP_ROWS = 40

#: zip 里的记号 -> (中文名, 下一步)。出现任意一个就拒绝读这份工作簿。
UNSUPPORTED_FEATURES = (
    ("xl/vbaProject.bin", "带宏（VBA）的工作簿",
     "在 Excel 里「另存为 → Excel 工作簿(.xlsx)」去掉宏，或另存为 CSV"),
    ("xl/externalLinks/", "带外部链接的工作簿",
     "在 Excel 里「数据 → 编辑链接 → 断开链接」，或另存为 CSV"),
    ("xl/charts/", "带图表的工作簿",
     "在 Excel 里删掉图表后另存为 .xlsx，或直接另存为 CSV"),
)


def _openpyxl():
    """把 import 失败说成人话（运行期依赖，规则 4 的 ADR-0003 就是为它写的）。"""
    try:
        import openpyxl
    except ImportError as exc:                       # pragma: no cover - 装机问题
        raise ValueError(
            "读 Excel 需要 openpyxl（只用于读，写出仍走仓库自带的那一份）。"
            "下一步：装一次 `pip install openpyxl`，或者把表另存为 CSV 再导入。"
        ) from exc
    return openpyxl


def _parts(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as zf:
            return zf.namelist()
    except zipfile.BadZipFile as exc:
        raise ValueError(
            f"{path.name}: 这不是一个 .xlsx（.xlsx 其实是一个 zip 包，这份打不开）。"
            "如果它本来是 .xls 或别的格式改了后缀，请用 Excel 另存为 .xlsx 或 CSV 再导入。"
        ) from exc


def _has_formulas(path: Path, parts: list[str]) -> bool:
    """表里有没有公式。只看每张工作表 XML 的开头一段——表头之后就是数据行，
    1 MB 足够盖住前几百行；真去扫整份会为一个大文件多解压一次几百 MB。"""
    with zipfile.ZipFile(path) as zf:
        for name in parts:
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                with zf.open(name) as handle:
                    if b"<f" in handle.read(1 << 20):
                        return True
    return False


def _reject_unsupported(path: Path) -> bool:
    """三种读不了的形态就地报错；返回"表里有没有公式"。"""
    parts = _parts(path)
    for marker, label, advice in UNSUPPORTED_FEATURES:
        if any(name == marker or name.startswith(marker) for name in parts):
            raise ValueError(
                f"{path.name}: 这是一份{label}，我们只读单元格里的数字与文本，"
                f"读不了它的宏 / 链接 / 图表。下一步：{advice}，再导入。"
            )
    return _has_formulas(path, parts)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:g}"
    return str(value)


def _top(ws, limit: int = TOP_ROWS) -> list[list[str]]:
    return [[_text(cell) for cell in row]
            for row in ws.iter_rows(min_row=1, max_row=limit, values_only=True)]


def _frame(ws, skip: int, width: int) -> pd.DataFrame:
    """数据区（表头之后）原样读进来：数字、文本、日期都保留，转换交给 csvlog。"""
    rows = [list(row) for row in
            ws.iter_rows(min_row=skip + 1, max_col=max(1, width), values_only=True)]
    return pd.DataFrame(rows)


def _uncached_formula(ws_formulas, ws_values, limit: int = 2000) -> bool:
    """有没有"有公式、但没存结果"的格子。

    Excel 自己存盘时会连结果一起存，所以读缓存值就够了；别的工具生成的工作簿常常
    只有公式、没有缓存值，读出来就是一片空白——那是"看着成功其实缺内容"。
    """
    for row_f, row_v in zip(
        ws_formulas.iter_rows(min_row=1, max_row=limit, values_only=True),
        ws_values.iter_rows(min_row=1, max_row=limit, values_only=True),
    ):
        for formula, value in zip(row_f, row_v):
            if isinstance(formula, str) and formula.startswith("=") and value is None:
                return True
    return False


def _load(openpyxl, path: Path, data_only: bool):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=data_only,
                                      keep_links=False)
    except Exception as exc:
        raise ValueError(
            f"{path.name}: 打不开这份 Excel（{type(exc).__name__}: {exc}）。"
            "下一步：用 Excel 打开它另存为 .xlsx 或 CSV 再试。"
        ) from exc


def sheet_names(path: str | Path) -> list[str]:
    """工作簿里有哪些 sheet（给 ``--sheet`` 和导入报告用）。"""
    path = Path(path)
    openpyxl = _openpyxl()
    book = _load(openpyxl, path, data_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_xlsx_session(
    path: str | Path,
    sheet: str | None = None,
    renames: dict[str, str] | None = None,
    units: dict[str, str] | None = None,
):
    """把一份 ``.xlsx`` 读成场次。

    ``sheet`` 不给时**按顺序试**，用第一张能当通道表读的：自己的导出把「元数据」
    放在第一张，所以不能按位置硬取第一张。试不成的 sheet 会把原因原样带进最后那句
    报错里——"哪张不行、为什么"比"读不出来"有用。
    """
    path = Path(path)
    has_formulas = _reject_unsupported(path)
    openpyxl = _openpyxl()
    chosen = str(sheet or "").strip() or csvlog.load_sheet(path)
    book = _load(openpyxl, path, data_only=True)
    formulas = _load(openpyxl, path, data_only=False) if has_formulas else None
    try:
        names = list(book.sheetnames)
        if not names:
            raise ValueError(f"{path.name}: 这份工作簿里一张 sheet 都没有。")
        if chosen:
            if chosen not in names:
                raise ValueError(
                    f"{path.name}: 没有叫 {chosen!r} 的 sheet。这份工作簿里有："
                    + "、".join(names)
                    + "。下一步：换成上面某一个名字，或去掉 --sheet 让工具自己挑。"
                )
            order = [chosen]
        else:
            order = names
        problems: list[str] = []
        for name in order:
            ws = book[name]
            top = _top(ws)
            if not top:
                problems.append(f"{name}: 是空的")
                continue
            try:
                _head, skip = csvlog.layout_of(top)
            except ValueError as exc:
                problems.append(f"{name}: {exc}")
                continue
            width = max(len(row) for row in top)
            frame = _frame(ws, skip, width)
            if frame.empty:
                problems.append(f"{name}: 表头之后没有数值行")
                continue
            if formulas is not None and _uncached_formula(formulas[name], ws):
                raise ValueError(
                    f"{path.name} 的「{name}」里有公式**没有存结果**（别的工具生成的表常这样），"
                    "读出来会是一片空白。下一步：用 Excel 打开另存一次（它会补上结果），"
                    "或者把这张表另存为 CSV 再导入。"
                )
            try:
                session = csvlog.session_from_frame(
                    path, top, frame, renames=renames, units=units,
                    fmt="xlsx", sheet=name,
                )
            except SidecarError:
                raise
            except ValueError as exc:
                problems.append(f"{name}: {exc}")
                continue
            session.header["sheets"] = names
            return session
        raise ValueError(
            f"{path.name}: 这份工作簿里没有能当通道表读的 sheet（"
            + "；".join(problems[:4])
            + "）。下一步：把要导入的那张表另存为 CSV，或在那张表里加上一列时间"
            "（列名用 Time / t / Timestamp）。"
        )
    finally:
        book.close()
        if formulas is not None:
            formulas.close()
